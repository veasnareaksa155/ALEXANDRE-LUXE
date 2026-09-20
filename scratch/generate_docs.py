import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_project_planning_docx(output_path):
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    GOLD = RGBColor(197, 160, 36)
    DARK = RGBColor(30, 30, 30)
    GRAY = RGBColor(100, 100, 100)

    # Title Banner
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("LEGACY STORE")
    title_run.font.name = "Georgia"
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = GOLD

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("E-Commerce Project Planning Sheet & Executive Specification")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(14)
    sub_run.font.italic = True
    sub_run.font.color.rgb = GRAY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.name = "Georgia"
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = DARK
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = "Georgia"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = GOLD
        return h

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(11)
            r_bold.font.bold = True
            r_bold.font.color.rgb = DARK
        r_text = p.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(11)
        r_text.font.color.rgb = DARK
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(11)
            r_bold.font.bold = True
            r_bold.font.color.rgb = DARK
        r_text = p.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(11)
        r_text.font.color.rgb = DARK
        return p

    # PILLAR 01: DEFINE BUSINESS
    add_heading_1("01 | កំណត់ Business (Define Business)")
    add_bullet(" Legacy Store — An ultra-luxury E-Commerce & Real-Time Delivery Satellite platform for high-end fashion, haute couture, leather goods, and premium accessories.", "1. Business Name & Idea:")
    add_bullet(" Luxury Designer Outfits (Suits, Dresses, Jackets), Handbags, Leather Goods, Accessories, combined with Express Courier Delivery with Live GPS Tracking.", "2. Products / Services:")
    add_bullet(" High-Net-Worth Individuals (HNWIs), VIP clients, fashion enthusiasts, and luxury gift shoppers looking for high craftsmanship and fast doorstep delivery.", "3. Target Customer:")
    add_bullet(" Desire for authentic luxury fashion without shipping delays, requirement for transparent real-time GPS tracking, and convenience of instant mobile payment (Bakong KHQR / Card).", "4. Customer Need:")
    add_bullet(" B2C Retail E-Commerce with premium product margins + VIP Concierge Membership offering private drop access and priority courier service.", "5. Business Model:")

    # PILLAR 02: DEFINE PROCESS
    add_heading_1("02 | កំណត់ Process (Define Process)")
    add_bullet(" Product catalog curation, order processing & quality inspection at Legacy Atelier, real-time express courier dispatch, and instant KHQR payment verification.", "1. Main Activities:")
    add_bullet(" Interactive Storefront, Smart Discovery Search Popover, KHQR Payment Gateway, Live Delivery Map Modal with moving courier icon, Mobile Driver Site (/delivery), SuperAdmin Dashboard (/admin).", "2. Main Features:")
    add_bullet(" Discovery & Search → Inspection & Reviews → Bag & Checkout (KHQR/Card) → Order Success → Real-Time GPS Tracking Modal → Signature Handover.", "3. Customer Journey:")
    add_bullet(" Step 1: Order Creation (Client checkout) → Step 2: Fulfillment (Atelier packages item & assigns courier) → Step 3: Delivery & Live GPS (Courier streams GPS location on driver app) → Step 4: Completion (Handover & status updated to Delivered).", "4. Basic Business Process:")

    # PILLAR 03: WEBSITE PLANNING
    add_heading_1("03 | Website Planning")
    add_heading_2("1. Basic Sitemap")
    add_bullet(" Hero Showcase, Category Quick-links, Featured Collections, Best Sellers, Footer.", "Home Page (/):")
    add_bullet(" Full Product Catalog, Category Filters, Price Sorting, Search Discovery Popover.", "Shop Catalog (/shop):")
    add_bullet(" Image Gallery, Size/Color Selectors, Rating Stars, Customer Reviews & Rating Form.", "Product Details Modal:")
    add_bullet(" Slide-over drawers for cart management and wishlist item saving.", "Shopping Bag & Wishlist Drawers:")
    add_bullet(" Customer Info Input, Payment Selector (Bakong KHQR / Card / COD), Price Summary.", "Checkout Modal:")
    add_bullet(" Interactive Route Map, Live Moving Truck, Courier Phone Call, Item Summary, Paid Receipt.", "Delivery Satellite Modal:")
    add_bullet(" Unique per-user account details, saved wishlist, personal order history.", "User Account Page (/account):")
    add_bullet(" Mobile login, active order selector, real-time drive GPS simulator.", "Mobile Courier Driver App (/delivery):")
    add_bullet(" Analytics overview, Product/Category CRUD, Order status manager, Driver credentials.", "SuperAdmin Panel (/admin):")

    add_heading_2("2. Basic User Flow")
    add_bullet(" Browse Store → Filter/Search Item → View Product Details → Select Size/Color → Add to Bag → Checkout with KHQR → Track Live GPS Delivery → Post Review.", "Client Flow:")
    add_bullet(" Open /delivery on Mobile → Login Driver Credentials → Select Assigned Order → Start Express Delivery Drive → Coordinates update live on customer map.", "Courier Driver Flow:")
    add_bullet(" Access /admin → Login SuperAdmin → View Sales Analytics → Update Order Delivery Status → Manage Catalog & Drivers.", "SuperAdmin Flow:")

    # PILLAR 04: OUTPUT
    add_heading_1("04 | Output (Project Planning Sheet & Presentation Outline)")
    add_heading_2("1. Executive Presentation Outline (Slide Deck Structure)")
    add_bullet(" Slide 1: Cover & Project Vision — Legacy Store Haute Couture E-Commerce")
    add_bullet(" Slide 2: Business & Target Market — Luxury retail market in Cambodia & HNWIs target segment")
    add_bullet(" Slide 3: Core Technology & Features — React 18, Laravel 11, Supabase, Bakong KHQR, Live Satellite GPS")
    add_bullet(" Slide 4: Website Sitemap & User Flow — Customer, Driver App, and Admin Dashboard architecture")
    add_bullet(" Slide 5: Real-Time Delivery Satellite Demonstration — Live route map & driver mobile app sync")
    add_bullet(" Slide 6: Results & Conclusion — Production status, data isolation, and deployment readiness")

    # Footer
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = footer_p.add_run("LEGACY STORE © 2026 — PROJECT PLANNING SHEET")
    f_run.font.name = "Calibri"
    f_run.font.size = Pt(9)
    f_run.font.color.rgb = GRAY

    doc.save(output_path)
    print(f"Planning sheet successfully generated at: {output_path}")

if __name__ == "__main__":
    project_root = r"d:\FULL_STACK_DEVELOPER\Alexandre Luxe"
    docx_filename = "Legacy_Store_Project_Planning_Sheet.docx"
    target_path = os.path.join(project_root, docx_filename)
    create_project_planning_docx(target_path)
    
    artifact_dir = r"C:\Users\MSI\.gemini\antigravity\brain\637bd956-0097-4a94-9304-c65feba3f3ff"
    if os.path.exists(artifact_dir):
        artifact_target = os.path.join(artifact_dir, docx_filename)
        create_project_planning_docx(artifact_target)
